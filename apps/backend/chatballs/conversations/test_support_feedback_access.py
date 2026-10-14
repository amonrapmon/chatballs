"""Current access applies identically to feedback totals, rows and native files."""

import csv
from io import BytesIO, StringIO

from django.db import connection, transaction
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.exceptions import PermissionDenied

from chatballs.conversations.models import Conversation
from chatballs.conversations.quality.criteria import configure_quality
from chatballs.conversations.services import delete_conversation
from chatballs.conversations.support_export import safe_csv_value
from chatballs.conversations.support_export_rows import report_tables
from chatballs.conversations.support_feedback_test_base import FeedbackTestBase
from chatballs.conversations.support_report import calculate_support_report
from chatballs.i18n import t
from chatballs.identity.group_models import EmployeeGroupMember
from chatballs.identity.models import OrganizationMembership
from chatballs.identity.supervision_models import EmployeeGroupSupervision
from chatballs.tenancy.context import TenantContext
from chatballs.tenancy.database import tenant_atomic


class FeedbackAccessTests(FeedbackTestBase):
    def prepared(self):
        self.reply()
        self.reply(user=self.colleague.user)
        event = self.complete()
        self.rated(event, comment=" =HYPERLINK(\"https://invalid.test\")")
        self.analyzed(event)

    def responses(self, **params):
        return [self.client.get(self.url.rstrip("/") + suffix, {**self.params, **params})
                for suffix in ("/", ".csv", ".xlsx")]

    def test_supervisor_redaction_current_scope_and_revocation_in_all_formats(self):
        self.prepared()
        OrganizationMembership.objects.filter(pk=self.idle.pk).update(role="SUPERVISOR")
        self.idle.refresh_from_db()
        link = EmployeeGroupSupervision.objects.create(organization=self.organization, employee=self.idle, group=self.group)
        EmployeeGroupMember.objects.create(organization=self.organization, employee=self.colleague, group=self.group)
        self.client.force_authenticate(self.idle.user)
        response = self.client.get(self.url, self.params)
        self.assertEqual(response.status_code, 200, response.data)
        for source in ("customerRatings", "qualityAnalysis"):
            self.assertEqual(response.data[source]["count"], 1)
            self.assertEqual(response.data[source]["details"][0]["responderIds"], [self.colleague.user_id])
        self.assertEqual([row["employee"] for row in response.data["employees"]], [self.colleague.user_id])
        self.assertTrue(all(response.status_code == 200 for response in self.responses()))
        # Ordinary assignment must not keep managerial feedback after revocation.
        Conversation.objects.filter(pk=self.conversation.pk).update(assigned_operator=self.idle.user)
        link.delete()
        report = self.client.get(self.url, self.params).data
        for source in ("customerRatings", "qualityAnalysis"):
            self.assertEqual((report[source]["count"], report[source]["details"]), (0, []))
        self.assertTrue(all(response.status_code == 404 for response in self.responses(group=self.group.pk)))

    def test_employee_stale_context_block_and_unavailable_dimensions(self):
        self.prepared()
        OrganizationMembership.objects.filter(pk=self.member.pk).update(role="EMPLOYEE")
        result = self.report(employee=self.user.pk)
        for source in ("customerRatings", "qualityAnalysis"):
            self.assertEqual(result[source]["state"], "forbidden")
            self.assertIsNone(result[source]["count"])
            self.assertEqual(result[source]["details"], [])
            self.assertEqual(result["employees"][0][source], result[source])
        responses = self.responses(employee=self.user.pk)
        self.assertTrue(all(response.status_code == 200 for response in responses))
        self.assertTrue(all(b"HYPERLINK" not in response.content for response in responses[1:]))
        OrganizationMembership.objects.filter(pk=self.member.pk).update(role="OWNER")
        for key, value in (("employee", self.foreign_user.pk), ("group", self.foreign_group.pk),
                           ("connection", self.foreign_connection.pk)):
            self.assertTrue(all(response.status_code == 404 for response in self.responses(**{key: value})))
        OrganizationMembership.objects.filter(pk=self.member.pk).update(blocked_at=timezone.now())
        with self.assertRaises(PermissionDenied):
            self.report()
        self.assertTrue(all(response.status_code == 404 for response in self.responses()))

    def test_runtime_tenant_isolation_and_physical_deletion(self):
        self.prepared()
        foreign_context = TenantContext.for_membership(self.foreign_member)
        with transaction.atomic():
            with connection.cursor() as cursor:
                cursor.execute("SET LOCAL ROLE chatballs_runtime_app")
            try:
                self.assertEqual(self.report()["customerRatings"]["count"], 1)
                foreign = calculate_support_report(context=foreign_context, filters=self.filters)
                self.assertEqual(foreign["customerRatings"]["count"], 0)
                self.assertEqual(foreign["qualityAnalysis"]["count"], 0)
                with tenant_atomic(self.context):
                    delete_conversation(context=self.context, conversation=self.conversation)
                self.assertEqual(self.report()["customerRatings"]["count"], 0)
                self.assertEqual(self.report()["qualityAnalysis"]["count"], 0)
            finally:
                with connection.cursor() as cursor:
                    cursor.execute("RESET ROLE")
        self.assertTrue(Conversation.objects.filter(pk=self.foreign_conversation.pk).exists())

    def test_removed_membership_and_disabled_features_keep_the_original_result_snapshot(self):
        self.prepared()
        self.colleague.delete()
        configure_quality(context=self.context, enabled=False)
        self.policy.enabled = False
        self.policy.save(update_fields=["enabled"])
        self.reply(user=self.idle.user)
        report = self.report()
        for source in ("customerRatings", "qualityAnalysis"):
            self.assertEqual(report[source]["count"], 1)
            self.assertEqual(report[source]["details"][0]["responderIds"], [self.user.pk])
            counts = {row["employee"]: row[source]["count"] for row in report["employees"]}
            self.assertEqual(counts, {self.user.pk: 1, self.idle.user_id: 0})

    def test_http_json_csv_xlsx_feedback_values_match_and_strings_cannot_be_formulas(self):
        self.prepared()
        responses = self.responses()
        self.assertTrue(all(response.status_code == 200 for response in responses))
        report = responses[0].json()
        tables = report_tables(report)
        csv_rows = list(csv.reader(StringIO(responses[1].content.decode("utf-8-sig"))))
        expected_csv = []
        for name, rows in tables:
            expected_csv.append([t(f"support_stats.sheet.{name}")])
            expected_csv.extend([[str(safe_csv_value(value)) for value in row] for row in rows])
            expected_csv.append([])
        generated = t("support_stats.generated")
        def clean(rows):
            return [row for row in rows if not row or row[0] != generated]
        self.assertEqual(clean(expected_csv), clean(csv_rows))
        workbook = load_workbook(BytesIO(responses[2].content), data_only=False)
        for sheet, (_, rows) in zip(workbook, tables, strict=True):
            expected = [[t("support_stats.unavailable") if value is None else None if value == "" else value
                         for value in row] for row in rows]
            self.assertEqual(clean(expected), clean([list(row) for row in sheet.values]))
            self.assertTrue(all(cell.data_type != "f" for row in sheet for cell in row))
        groups = workbook[t("support_stats.sheet.rating_groups")]
        self.assertEqual(groups.cell(2, 5).data_type, "n")
        self.assertEqual(groups.cell(2, 7).data_type, "n")
        self.assertEqual(report["customerRatings"]["count"], 1)
        self.assertEqual(report["qualityAnalysis"]["count"], 1)
